from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SRC = Path("output/patent_disclosure/slmf_bbdm_patent_innovation_integrated.docx")
OUT = Path("output/patent_disclosure/slmf_bbdm_patent_two_innovations_with_losses.docx")


def set_east_asian_font(run, font="SimSun") -> None:
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font)


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        for run in paragraph.runs:
            run.text = ""
        first.text = text
        set_east_asian_font(first)
    else:
        run = paragraph.add_run(text)
        set_east_asian_font(run)


def insert_after(paragraph, text: str, style_name: str = "Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    paragraph._p.getparent().remove(new_p)
    new_para.style = style_name
    run = new_para.add_run(text)
    set_east_asian_font(run)
    return new_para


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def find_para(doc: Document, exact: str):
    for p in doc.paragraphs:
        if p.text.strip() == exact:
            return p
    raise ValueError(f"paragraph not found: {exact}")


def main() -> None:
    doc = Document(SRC)

    # Reframe the two innovation headings so loss functions are explicitly part of them.
    replace_paragraph_text(
        find_para(doc, "创新点一：面向盆腔小病灶的频带分解桥式扩散状态构造机制"),
        "创新点一：面向盆腔小病灶的频带分解桥式扩散状态与高摄取-频域联合损失机制",
    )
    replace_paragraph_text(
        find_para(doc, "创新点二：盆腔器官代谢规则驱动的时间门控轻量适配与 SUV 误差重分配机制"),
        "创新点二：盆腔器官代谢规则驱动的时间门控适配与 SUV-冷区假热点阶段化损失机制",
    )

    # Add small-lesion/frequency/spatial-binding losses into innovation point one.
    anchor_one = find_para(
        doc,
        "对于 PNG 或其他二维图像格式，也可采用文件名匹配方式配对 CT、PET 和标签图，并统一缩放到模型输入大小；若需要严格 SUV 损失和物理尺度分析，则优选 DICOM 到 NPZ 缓存流程。由此，图像输入、物理标定和配准质控被并入创新点一的桥式状态构造链条，而不是作为独立技术方案列出。",
    )
    innovation_one_loss = [
        "进一步地，创新点一还包括与频带桥式状态相配套的小病灶高摄取-频域联合损失。该损失不是普通像素损失的附加项，而是与 ρ_b(p)、A_hot(p) 和桥式时间步 τ 共同作用：在晚期去噪阶段，提高真实 PET 高摄取区域、热点候选区域和高频边缘区域的误差权重，使模型在恢复整体 PET 分布的同时避免小病灶峰值被背景平均化。",
        "基础重建项可写为 L_base=w_mse(τ)||ŷ_0-x_0||_2^2+w_l1||ŷ_0-x_0||_1+w_g(τ)||∇ŷ_0-∇x_0||_1，其中 w_mse(τ) 与 w_g(τ) 随桥式时间步变化，早期偏重全局强度，晚期偏重边缘梯度。由此，损失函数与桥式扩散阶段形成绑定，而不是独立于生成过程的常规重建误差。",
        "小病灶高摄取项可写为 M_top=1{x_0>=TopK_threshold(x_0,k)}，L_top=Σ_p M_top(p)·(1-exp(-|e_p|))^γ·|e_p|/(Σ_p M_top(p)+ε)。其中 M_top 可由真实 PET 高摄取区域、病灶掩膜或热点候选图共同确定，使极少比例的小病灶像素不会被大面积背景稀释。",
        "频域恢复项可写为 L_freq=mean(|F(ŷ_0)-F(x_0)|^α·|F(ŷ_0)-F(x_0)|)，并可叠加空间绑定项 L_NCE=-log exp(sim(z_i^CT,z_i^PET)/τ_n)/Σ_j exp(sim(z_i^CT,z_j^PET)/τ_n)。上述损失与频带桥式噪声共同构成创新点一：前向过程保护高频和热点候选区域，反向训练过程再对这些区域给予更高恢复权重。",
    ]
    last = anchor_one
    for text in innovation_one_loss:
        last = insert_after(last, text)

    # Add SUV/organ/false-hotspot/uncertainty losses into innovation point two.
    anchor_two = find_para(
        doc,
        "其中 L_i 可包括小病灶高摄取损失、频域恢复损失、空间绑定损失、SUV 定量损失、器官代谢规则损失、冷区假热点抑制损失、热点候选监督损失和异方差不确定性损失；g_i(τ,v_i) 为同时受时间步 τ 和临床有效性标志 v_i 控制的门控函数。通过该设计，模型不会把优化能力平均分配给大面积背景，而是把误差权重转移到病灶高摄取区域、边界过渡区域、临床定量敏感区域以及容易产生假热点的冷区。",
    )
    innovation_two_loss = [
        "创新点二中的损失函数重点体现盆腔器官代谢规则和临床定量约束。ROI-SUV 损失先将归一化 PET 反变换到 SUV 物理空间，再在病灶掩膜或热点候选区域约束 SUVmax、SUVmean 和肿瘤背景比 TBR：L_SUV=|SUVmax_pred-SUVmax_gt|+|SUVmean_pred-SUVmean_gt|+0.5·|TBR_pred-TBR_gt|。当样本缺少可靠 SUV 标定字段时，该损失由临床有效性标志 v_SUV 自动屏蔽。",
        "器官代谢规则损失用于把不同盆腔器官区别对待，可写为 L_organ=Σ_c η_c·mean(ReLU(ŷ_0)·M_cold^c)。其中骨、脂肪和肌肉等冷区惩罚异常正向高值，膀胱和直肠或肠道允许生理摄取，子宫或盆腔区域采用弱约束或候选增强约束。该损失使器官分区不只是条件输入，而是直接参与训练目标。",
        "冷区假热点抑制损失可写为 L_false=mean(|ŷ_0(p)|·1{ŷ_0(p)>Q_q(ŷ_0⊙M_cold)}·M_cold(p))，用于在非病灶且非生理高摄取区域内寻找高分位异常响应并进行惩罚。该项与器官分区规则绑定，区别于普通背景抑制损失。",
        "异方差不确定性损失可写为 L_NLL=0.5·exp(-s)·(ŷ_0-x_0)^2+0.5·s，s=clip(logσ_y^2)。最终，创新点二的阶段化目标可统一写为 L_total=L_base+Σ_i λ_i·g_i(τ,v_i)·L_i，其中 g_i(τ,v_i) 同时受桥式时间步和临床有效性标志控制，使 SUV、器官、冷区假热点和不确定性约束仅在适当样本和适当去噪阶段发挥作用。",
    ]
    last = anchor_two
    for text in innovation_two_loss:
        last = insert_after(last, text)

    # Remove the standalone loss-function section; its content has been integrated above.
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "损失函数与公式化描述" and p.style.name.startswith("Heading"):
            start = i
        elif start is not None and p.text.strip() == "有益效果" and p.style.name.startswith("Heading"):
            end = i
            break
    if start is not None and end is not None:
        for p in list(doc.paragraphs[start:end]):
            remove_paragraph(p)

    # Sync downstream wording with the new two-innovation framing.
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("与传统 CT-to-PET 合成方法相比"):
            replace_paragraph_text(
                p,
                "与传统 CT-to-PET 合成方法相比，本发明至少具有以下有益效果。第一，本发明不是泛化地采用布朗桥扩散，而是在桥式状态中引入频带同步分解、空间门控噪声以及与之配套的小病灶高摄取-频域联合损失，降低普通跨模态扩散的结构漂移和小病灶平滑风险。第二，本发明不是单纯使用边缘滤波或 Gabor 响应，而是将方向频率响应与热点候选、Top-K 高摄取损失、频域恢复损失共同用于高频保护。第三，本发明不是简单复制通用控制分支，而是用浅层/深层分工的零初始化尺度适配单元分别注入频率热点信息和器官代谢信息。第四，本发明不是只追求视觉相似，而是把 ROI-SUV、器官代谢规则、冷区假热点和异方差不确定性纳入同一阶段化损失机制。",
            )
        elif txt.startswith("图4 为核心创新点示意图"):
            replace_paragraph_text(
                p,
                "图4 为核心创新点示意图，展示频带桥式状态构造、小病灶高摄取-频域联合损失、盆腔器官代谢规则注入和 SUV/冷区假热点阶段化损失之间的关系。",
            )
        elif txt.startswith("1. 一种面向子宫内膜癌 CT 图像的 PET 合成方法"):
            replace_paragraph_text(
                p,
                "1. 一种面向子宫内膜癌 CT 图像的 PET 合成方法，其特征在于，包括：获取配准后的 CT 图像，并在训练阶段获取对应 PET 图像；根据所述 CT 图像构建方向频率响应能量图、热点候选概率图以及盆腔器官分区特征；对所述 CT 图像、PET 图像和噪声进行对应频带分解，并在每一频带内根据桥式混合系数构造扩散状态；其中每一频带的噪声强度由频带倍率、方向频率响应能量、热点候选概率和盆腔器官分区规则共同确定；训练时进一步采用与所述扩散状态相绑定的小病灶高摄取-频域联合损失以及与所述盆腔器官分区规则相绑定的 SUV/冷区假热点阶段化损失；将所述扩散状态和所述盆腔器官分区特征输入时间门控的条件去噪网络，输出合成 PET 图像。",
            )
        elif txt.startswith("6. 根据权利要求1所述的方法"):
            replace_paragraph_text(
                p,
                "6. 根据权利要求1所述的方法，其中训练损失包括小病灶高摄取损失、频域恢复损失、空间绑定损失、ROI-SUV 代谢定量损失、冷区假热点抑制损失、器官代谢规则损失和异方差不确定性损失中的至少三种；其中小病灶高摄取损失、频域恢复损失与频带桥式扩散状态共同构成第一创新约束，ROI-SUV 代谢定量损失、冷区假热点抑制损失和器官代谢规则损失与盆腔器官分区共同构成第二创新约束。",
            )
        elif txt.startswith("从权利布局角度看"):
            replace_paragraph_text(
                p,
                "从权利布局角度看，可围绕方法、系统、装置、介质和模型训练方法分别撰写权利要求。独立权利要求宜覆盖两组组合技术特征：其一为“CT 端点约束的频带桥式扩散状态、方向频率响应与热点候选共同调制高频噪声、小病灶高摄取-频域联合损失”；其二为“盆腔器官代谢规则、时间门控尺度适配、ROI-SUV/冷区假热点/器官规则阶段化损失”。从属权利要求再限定 Gabor 作为方向频率响应的一种实现、六类器官规则、ROI-SUV 有效性屏蔽、不确定性输出和三维扩展。这样既能保护当前工程实现，也能减少被单个公开模块覆盖的风险。",
            )
        elif txt.startswith("本发明围绕子宫内膜癌 CT-to-PET 合成"):
            replace_paragraph_text(
                p,
                "本发明围绕子宫内膜癌 CT-to-PET 合成中的小病灶易平滑、代谢热点易错位、器官生理摄取易混淆和临床 SUV 指标难保持等问题，重构形成两个主要创新点。第一创新点是频带分解桥式扩散状态与小病灶高摄取-频域联合损失的协同机制；第二创新点是盆腔器官代谢规则驱动的时间门控适配与 SUV/冷区假热点阶段化损失机制。该方案的可保护重点不是普通布朗桥扩散、Gabor、U-Net、零卷积或单个损失函数本身，而是桥式状态、条件适配和创新损失之间的限定关系。由此，模型能够在理论结构和工程实现上兼顾生成质量、医学可解释性和临床定量约束，适合作为专利申请中的方法、系统、设备和存储介质方案进行布局。",
            )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
