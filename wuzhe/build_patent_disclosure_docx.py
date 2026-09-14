from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("output/patent_disclosure")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_DOCX = OUT_DIR / "SLMF_BBDM_子宫内膜癌CT_to_PET合成专利交底书.docx"

FIGURES = [
    (Path("output/figures/slmf_bbdm_png256_architecture.png"), "图1  SLMF-BBDM 总体网络结构示意图"),
    (Path("output/midterm_report/figures/fig1_slmf_bbdm_pipeline_sci.png"), "图2  CT-to-PET 布朗桥扩散训练与推理流程图"),
    (Path("output/midterm_report/figures/fig4_innovation_points.png"), "图3  两项核心创新点与技术闭环示意图"),
    (Path("output/midterm_report/figures/fig2_data_registration_qc.png"), "图4  PET/CT 数据配准与质控流程示意图"),
]


NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(88, 88, 88)
LIGHT = "F4F6F9"
HEADER = "E8EEF5"


def set_run_font(run, size=None, bold=None, color=None, name="SimSun"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level < 3 else NAVY)
    return p


def add_para(doc, text, bold_prefix=None, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.3)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=10.5, name="Cambria Math")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=GRAY)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    doc.add_paragraph()


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, width in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = Inches(width)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, HEADER)
                set_cell_text(cell, text, bold=True, color=NAVY, size=10)
            else:
                set_cell_text(cell, text, size=9.5)
    doc.add_paragraph()
    return table


def add_figure(doc, path, caption, width=6.2):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rh = header.add_run("专利技术交底书 | SLMF-BBDM 子宫内膜癌 CT-to-PET 合成")
    set_run_font(rh, size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = footer.add_run("内部技术资料")
    set_run_font(rf, size=9, color=GRAY)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("专利技术交底书")
    set_run_font(r, size=24, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("一种面向子宫内膜癌 CT 图像的代谢保真 PET 合成方法、系统、设备及存储介质")
    set_run_font(r, size=15, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于小病灶频率保护布朗桥扩散与医学先验轻量注入的 SLMF-BBDM 技术方案")
    set_run_font(r, size=12, color=GRAY)

    add_table(doc, [
        ["项目", "填写内容"],
        ["发明名称", "一种面向子宫内膜癌 CT 图像的代谢保真 PET 合成方法、系统、设备及存储介质"],
        ["技术方向", "医学图像智能生成；CT-to-PET 跨模态合成；扩散模型；妇科肿瘤影像分析"],
        ["核心模型", "SLMF-BBDM：Small-Lesion Metabolic-Fidelity Brownian Bridge Diffusion Model"],
        ["交底重点", "小病灶高频结构保护、自适应布朗桥噪声、医学先验轻量注入、SUV 代谢保真与假热点抑制"],
        ["适用对象", "子宫内膜癌及盆腔肿瘤 PET/CT 检查中 CT 图像到 PET 代谢影像的辅助合成与分析"],
    ], [1.5, 4.8])

    add_callout(
        doc,
        "交底说明",
        "本交底书依据当前工作目录中的 SLMF-BBDM 模型实现、配置文件和阶段报告整理而成。文中实验数值不作未验证性能承诺，重点描述可被工程实施和权利要求抽取的技术结构、处理流程、公式关系及可替代实施方式。"
    )

    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_para(doc, "本发明涉及医学图像智能生成和妇科肿瘤影像辅助分析领域，具体提出一种面向子宫内膜癌 CT 图像的代谢保真 PET 合成方法、系统、设备及存储介质。该方法将源 CT 图像与目标 PET 图像构造为条件布朗桥扩散过程的两个端点，在训练阶段由真实 PET、源 CT 和随机噪声生成中间扩散状态，在推理阶段仅以 CT 及其派生先验为条件，经少步反向采样得到合成 PET 图像。为解决子宫内膜癌小病灶在 CT-to-PET 合成中容易被平滑、代谢热点位置易偏移、膀胱或肠道等生理摄取区域易产生假阳性的问题，本发明进一步引入多尺度频率保护、自适应噪声分配、Gabor 高频能量先验、器官结构先验、热点候选先验以及轻量级 Zero-Conv 条件注入机制。")
    add_para(doc, "在扩散噪声设计上，本发明将图像分解为拉普拉斯金字塔的低频、中频和高频分量，并根据不同频带设置不同噪声强度；其中高频噪声由 Gabor 能量图进一步调制，使小病灶边缘和局部代谢热点在扩散破坏过程中获得更强保护。在条件注入上，本发明不复制完整 ControlNet 分支，而是采用逐尺度 1×1 零初始化卷积将 CT 多尺度特征、Gabor 特征、热点概率图和器官距离特征注入 U-Net 跳跃连接，并以时间相关系数控制不同去噪阶段的先验贡献。")
    add_para(doc, "在优化目标上，本发明将基础重建损失、Top-K 小病灶聚焦损失、焦点频域损失、ROI-SUV 代谢定量损失、假热点抑制损失、器官一致性损失、PatchNCE 空间绑定损失以及异方差不确定性似然损失组合为阶段化训练目标，使模型在全局结构、局部边缘、临床 SUV 指标和器官代谢合理性之间形成协同约束。该技术方案可用于 PET 缺失或 PET 采集受限场景下的代谢影像辅助合成、病灶候选区域提示、术前影像评估和多模态影像研究。")

    add_heading(doc, "1. 技术领域", 1)
    add_para(doc, "本发明属于医学图像处理、深度学习生成模型和肿瘤影像辅助诊断技术领域，尤其涉及一种基于 CT 图像生成对应 PET 代谢图像的方法。更具体地，本发明面向子宫内膜癌及其盆腔相关病灶，利用布朗桥扩散模型建立 CT 解剖结构与 PET 代谢摄取之间的跨模态映射，并通过小病灶频率保护、医学先验注入和代谢定量约束提高合成 PET 图像对微小高摄取区域、器官边界和临床 SUV 指标的保持能力。")
    add_para(doc, "本发明既可作为独立的 CT-to-PET 合成模型使用，也可作为后续分割、检出、病灶热区提示、治疗前评估、影像组学特征提取或多中心影像数据补全的前置模块使用。其输出形式包括但不限于合成 PET 图像、热点先验图、代谢置信度图、不确定性图以及面向临床定量分析的 SUV 相关指标。")

    add_heading(doc, "2. 背景技术", 1)
    add_para(doc, "子宫内膜癌是常见妇科恶性肿瘤之一，影像检查在术前分期、病灶定位、淋巴结转移评估和治疗方案制定中具有重要价值。CT 图像能够提供较稳定的解剖结构信息，但对微小病灶、肿瘤代谢活性和局部高摄取区域的表达能力有限。PET 图像能够反映葡萄糖代谢和肿瘤活性，尤其在高代谢病灶识别方面具有优势，但 PET/CT 检查成本较高、采集流程复杂、辐射剂量更高，并且不同中心之间存在采集协议和 SUV 标定差异。因此，在仅有 CT 或 PET 不完整的场景下，如何生成具有临床可解释性的 PET 代谢图像，是医学图像智能生成领域具有现实意义的问题。")
    add_para(doc, "现有 CT-to-PET 合成方法通常采用 U-Net、GAN、条件 VAE 或普通扩散模型进行跨模态生成。这些方法能够在整体灰度分布上学习 CT 与 PET 的统计对应关系，但对子宫内膜癌场景仍存在若干不足。第一，盆腔病灶往往体积较小、边界不规则，PET 热点在二维切片中可能仅占极少比例，常规像素均值损失会使模型更倾向于优化大面积背景，从而导致小病灶被平滑。第二，膀胱、肠道等组织可出现生理性高摄取，若模型缺乏器官先验，则容易把正常生理摄取误解释为病灶热点，或者在骨、脂肪、肌肉等代谢冷区生成假热点。第三，普通扩散模型在加噪过程中通常对全部空间和频率成分采用统一噪声强度，高频边缘和微小热点更容易被破坏，反向采样时难以恢复。第四，已有生成模型常以视觉相似度为目标，缺少 SUVmax、SUVmean、肿瘤背景比等临床定量约束，导致合成结果虽然外观接近，却不一定满足代谢分析需求。")
    add_para(doc, "布朗桥扩散模型能够将源域图像和目标域图像作为桥的两端，通过源图像约束扩散路径，使反向生成过程更适合图像到图像翻译任务。然而，直接将布朗桥扩散应用于 CT-to-PET 合成仍不能自动解决小病灶保真、频带保护和器官代谢合理性问题。尤其在子宫内膜癌 PET/CT 数据中，病灶附近存在膀胱、直肠、骨盆等复杂结构，模型需要同时学习解剖约束、代谢热点约束和频率细节约束。因此，有必要提出一种将布朗桥扩散、自适应频率噪声、医学结构先验和临床定量损失统一起来的技术方案。")

    add_heading(doc, "3. 发明内容", 1)
    add_para(doc, "本发明的目的在于提供一种面向子宫内膜癌 CT 图像的代谢保真 PET 合成方法，使模型在仅输入 CT 图像及其可派生结构信息的条件下，生成既具有整体 PET 分布合理性、又能保持微小病灶高频边界和 SUV 定量特征的 PET 图像。")
    add_para(doc, "为实现上述目的，本发明采用 SLMF-BBDM 模型框架。该框架包括数据预处理模块、CT 多尺度编码模块、医学先验构建模块、尺度自适应布朗桥噪声模块、轻量条件注入模块、异方差 U-Net 去噪模块、阶段化代谢保真损失模块以及少步反向采样模块。模型训练时，输入配准后的 CT 图像、真实 PET 图像、病灶掩膜、可选器官掩膜、器官距离图、CT 衰减映射、PET SUV 张量和病例元数据；模型推理时，至少输入 CT 图像即可生成合成 PET，若具备器官或元数据信息则进一步提升约束强度。")
    add_callout(doc, "本发明拟合出的两个核心创新点", "创新点一：小病灶频率保护的尺度自适应布朗桥扩散机制。创新点二：面向临床代谢保真的多源医学先验轻量注入与阶段化误差重分配机制。")

    add_heading(doc, "4. 两个创新点", 1)
    add_heading(doc, "4.1 创新点一：小病灶频率保护的尺度自适应布朗桥扩散机制", 2)
    add_para(doc, "现有扩散模型在前向加噪时通常对整幅图像采用统一噪声系数，导致局部高频结构与小病灶热区在早期被过度扰动。本发明将布朗桥扩散过程与拉普拉斯金字塔频带分解结合，针对低频背景、中频器官轮廓和高频病灶边缘分别设置噪声强度，并引入 Gabor 高频能量图对高频噪声进行空间调制。由此，模型在保持布朗桥 CT 端点约束的同时，对小病灶边缘、热点峰值和局部纹理实施保护。")
    add_para(doc, "具体而言，设真实 PET 图像为 x0，源 CT 图像为 xs，时间步归一化变量为 τ=t/T，布朗桥混合系数为 m_t。常规桥式扩散可写为：")
    add_formula(doc, "x_t = m_t · x_s + (1 - m_t) · x_0 + σ_t · ε")
    add_para(doc, "本发明将噪声 ε 分解为多尺度残差 {ε_h, ε_m, ε_l}，并将真实 PET 和源 CT 同步分解为对应频带。对任一频带 b，前向状态表示为：")
    add_formula(doc, "x_t^b = m_t · x_s^b + (1 - m_t) · x_0^b + σ_t^b(E_g) · ε^b")
    add_para(doc, "其中 b 表示高频、中频或低频分量，E_g 表示由可学习 Gabor 滤波器组从 CT 提取的高频能量图。低频噪声倍率可设为 1.0，中频噪声倍率可设为 0.75，高频噪声倍率可设为 0.45，并在高频区域采用如下调制：")
    add_formula(doc, "σ_t^high(p) = σ_t · λ_high · [1 - α · clip(E_g(p), 0, 1)]")
    add_para(doc, "其中 p 为空间位置，λ_high 为高频噪声倍率，α 为 Gabor 调制强度。该式使边缘和疑似病灶高响应区域获得更低噪声破坏，而平坦区域仍保持足够随机扰动，从而避免模型仅记忆局部纹理。")
    add_heading(doc, "4.2 创新点二：面向临床代谢保真的多源医学先验轻量注入与阶段化误差重分配机制", 2)
    add_para(doc, "本发明并不将全部先验简单拼接到输入层，而是将 CT 多尺度特征、Gabor 高频特征、器官掩膜与距离图、CT 衰减映射、热点候选图、可选语义 token 和病例元数据组织为条件包，并通过逐尺度零初始化 1×1 卷积注入 U-Net 跳跃连接。该结构具有两个优点：一是零初始化卷积在训练起点等价于恒等扰动，不破坏基础扩散主干；二是每个尺度仅增加少量参数，相比复制完整 ControlNet 分支更轻量，适合医学小样本场景。")
    add_para(doc, "第 l 个尺度的跳跃连接可表示为：")
    add_formula(doc, "h_l' = h_l + β_l(τ) · Z_l([f_l^CT, f_l^organ, f^Gabor, A^hotspot])")
    add_para(doc, "其中 h_l 为 U-Net 原始跳跃特征，Z_l 为第 l 层零初始化卷积，β_l(τ) 为随扩散时间变化的注入系数。浅层在晚期去噪阶段更强调高频边缘和热点位置，深层在早中期更强调全局解剖和器官结构。")
    add_para(doc, "此外，本发明将误差重分配设计为阶段化损失堆栈。早期去噪关注器官级代谢分布和空间绑定，中期关注热点候选与频域恢复，晚期关注小病灶 Top-K 像素、SUV 定量和假热点抑制。总体损失可写为：")
    add_formula(doc, "L = L_base + Σ_i w_i · g_i(τ) · L_i")
    add_para(doc, "其中 L_i 可包括 Top-K 病灶损失、焦点频域损失、PatchNCE 损失、ROI-SUV 损失、假热点抑制损失、器官一致性损失、热点先验监督损失和异方差负对数似然损失，g_i(τ) 为与扩散阶段相关的平滑门控函数。通过该设计，模型不会把全部学习能力平均分配给大面积背景，而会把优化重点重分配到病灶高摄取区域、边界过渡区域和临床定量敏感区域。")

    for fig, cap in FIGURES[:3]:
        add_figure(doc, fig, cap)

    add_heading(doc, "5. 技术方案", 1)
    add_heading(doc, "5.1 数据预处理与输入构建", 2)
    add_para(doc, "在一种实施方式中，首先收集同一患者、同一或近似层面的 CT 与 PET 图像，并进行空间配准、切片匹配和质量控制。对于 DICOM 数据，读取 CT 像素并转换为 HU 值，将 HU 裁剪到预设窗口后归一化到 [-1,1]；读取 PET 像素并在具备体重、注射剂量、采集时间和放射性核素半衰期等字段时计算 SUV，再按预设 SUV 上限归一化到 [-1,1]。")
    add_formula(doc, "CT_norm = 2 · clip((HU - HU_min)/(HU_max - HU_min), 0, 1) - 1")
    add_formula(doc, "PET_norm = 2 · clip(SUV/SUV_max, 0, 1) - 1")
    add_para(doc, "同时，依据 CT HU 构造 511keV 衰减映射 μ_map，可采用线性近似 μ=0.096×HU/1000+0.096。若具备器官分割工具，可对 CT 体数据执行器官分割，将子宫/盆腔、膀胱、直肠/肠道、骨、脂肪、肌肉或其他组织映射为固定六类器官通道，并计算每类器官的有符号距离变换。")
    add_para(doc, "对于 PNG 或其他二维图像格式，也可采用文件名匹配方式配对 CT、PET 和标签图，并统一缩放到模型输入大小，例如 192×192 或 256×256。该实施方式适合早期验证和二维切片实验；若需要严格 SUV 损失和物理尺度分析，优选 DICOM 到 NPZ 缓存流程。")
    add_figure(doc, FIGURES[3][0], FIGURES[3][1])

    add_heading(doc, "5.2 CT 多尺度编码模块", 2)
    add_para(doc, "CT 多尺度编码模块用于从源 CT 中提取不同感受野的解剖结构特征。该模块包含 3×3 与 7×7 并行卷积的多尺度 stem，用于兼顾局部边缘和较大范围的组织结构；之后通过三级下采样获得四个尺度的特征图，通道数可分别为 64、128、256 和 256。得到的 CT 特征既作为扩散主干的结构条件，也作为 Zero-Conv 条件注入的基础。")
    add_formula(doc, "F_CT = {f_0^CT, f_1^CT, f_2^CT, f_3^CT} = Enc_CT(x_s)")
    add_para(doc, "与直接将 CT 和噪声 PET 在输入通道拼接相比，多尺度编码能够在浅层保留边缘和纹理，在深层保留盆腔器官布局和整体空间上下文。该设计有利于在 PET 合成中维持解剖一致性。")

    add_heading(doc, "5.3 高频 Gabor 先验模块", 2)
    add_para(doc, "Gabor 先验模块包含一组可学习 Gabor 滤波器。每个滤波器具有频率、方向、尺度、长宽比和相位等可学习参数，并通过约束变换保持物理意义。对 CT 图像进行卷积后得到多方向高频响应，再对各滤波器响应平方求和并开方，形成归一化 Gabor 能量图。")
    add_formula(doc, "G_k(x,y)=exp(-(x_θ^2+γ_k^2 y_θ^2)/(2σ_k^2)) · cos(2π f_k x_θ + φ_k)")
    add_formula(doc, "E_g = Norm( sqrt(Σ_k (G_k * x_s)^2) )")
    add_para(doc, "Gabor 特征在本发明中具有三重作用：其一，作为浅层 Zero-Conv 注入特征，帮助 U-Net 恢复小病灶边缘；其二，作为尺度自适应噪声调制因子，降低高频区域的噪声破坏；其三，配合焦点频域损失，缓解 PET 合成中的模糊现象。")

    add_heading(doc, "5.4 器官先验与热点先验模块", 2)
    add_para(doc, "器官先验模块将六类器官掩膜、六类器官距离图以及 μ_map 拼接为十三通道输入，通过三层轻量卷积生成三个尺度的器官特征。器官先验使模型能够区分病灶可能发生区域、生理高摄取区域和代谢冷区。例如膀胱与肠道允许较高生理摄取，骨、脂肪和肌肉区域则应抑制无依据热点。")
    add_para(doc, "热点先验模块是一个轻量 U-Net，输入 CT 与 Gabor 能量图，输出候选热点概率图 A^hotspot。训练阶段可由病灶掩膜、PET 高摄取分位图和可选距离图进行监督；推理阶段该热点图作为条件注意力引导扩散模型把有限生成能力聚焦到疑似病灶区域。")
    add_formula(doc, "A^hotspot = H_θ([x_s, E_g])")
    add_para(doc, "热点监督目标可由真实病灶掩膜 M 与 PET 高摄取分位区域 Q(PET) 取并集构造：")
    add_formula(doc, "Y_hot = max(M, 1{PET >= Quantile(PET, q)})")

    add_heading(doc, "5.5 轻量 Zero-Conv 条件注入模块", 2)
    add_para(doc, "本发明采用 Zero-Conv 条件注入替代复制式控制网络。对于第 0 层浅层尺度，条件由 CT 浅层特征、Gabor 特征和热点图拼接；对于第 1 至第 3 层，条件由 CT 对应尺度特征与器官先验特征拼接。每个尺度通过 1×1 零初始化卷积生成与 U-Net 跳跃连接同通道数的残差注入量。")
    add_formula(doc, "C_0 = concat(f_0^CT, f^Gabor, A^hotspot)")
    add_formula(doc, "C_l = concat(f_l^CT, f_l^organ), l∈{1,2,3}")
    add_formula(doc, "Δh_l = β_l(τ) · Conv1×1_zero(C_l)")
    add_para(doc, "由于 Conv1×1_zero 的参数在初始化时全为零，模型初始状态等价于未注入先验的基础扩散模型，训练过程中再逐渐学习每类先验对各尺度的贡献。这降低了先验错误或早期训练不稳定对主干网络的干扰。")

    add_heading(doc, "5.6 异方差 U-Net 去噪模块", 2)
    add_para(doc, "去噪主干采用残差 U-Net。其输入为当前扩散状态 x_t 与源 CT x_s 的拼接，若启用自条件机制，则额外拼接上一轮预测的 x0 估计。U-Net 编码器与解码器均包含时间嵌入调制的残差块，瓶颈处可接入语义 token 的交叉注意力，病例元数据可通过 FiLM 方式注入时间嵌入。输出通道包括预测 PET 图像和可选 log-variance 通道。")
    add_formula(doc, "ŷ_0, logσ_y^2 = UNet_θ([x_t, x_s, ŝ_0], t, Tokens, Meta, Δh_l)")
    add_para(doc, "异方差输出使模型不仅给出合成 PET，还能给出像素级观测不确定性。结合多次 Monte Carlo 采样，可进一步得到认知不确定性、偶然不确定性和综合置信度图。")
    add_formula(doc, "Var_total = Var_epistemic + exp(logσ_y^2)")
    add_formula(doc, "Conf = 1 - Var_total / max(Var_total)")

    add_heading(doc, "6. 损失函数与公式化描述", 1)
    add_para(doc, "本发明的训练目标不是单一像素误差，而是由多个与扩散阶段相关的损失项共同构成。基础重建损失包括 MSE、L1 和图像梯度损失，并根据 τ 调整权重：早期强调全局强度重建，晚期提高梯度和细节约束。")
    add_formula(doc, "L_base = w_mse(τ)||ŷ_0-x_0||_2^2 + w_l1||ŷ_0-x_0||_1 + w_g(τ)||∇ŷ_0-∇x_0||_1")
    add_para(doc, "Top-K 病灶聚焦损失只选择真实 PET 中亮度最高的 k% 像素参与计算，并引入 focal 权重，使难预测的高摄取像素产生更大梯度。")
    add_formula(doc, "M_top = 1{x_0 >= TopK_threshold(x_0,k)}")
    add_formula(doc, "L_topk = Σ_p M_top(p) · (1-exp(-|e_p|))^γ · |e_p| / (Σ_p M_top(p)+ε)")
    add_para(doc, "焦点频域损失在傅里叶域对预测 PET 与真实 PET 的频谱差异加权，频域误差越大，该频率分量权重越高。")
    add_formula(doc, "L_freq = mean( |F(ŷ_0)-F(x_0)|^α · |F(ŷ_0)-F(x_0)| )")
    add_para(doc, "ROI-SUV 损失先将归一化 PET 反变换到 SUV 物理空间，然后在病灶掩膜内约束 SUVmax、SUVmean 和肿瘤背景比 TBR。")
    add_formula(doc, "SUV(ŷ)=((ŷ+1)/2)·SUV_max")
    add_formula(doc, "L_SUV = |SUVmax_pred-SUVmax_gt| + |SUVmean_pred-SUVmean_gt| + 0.5·|TBR_pred-TBR_gt|")
    add_para(doc, "器官一致性损失用于抑制代谢冷区的异常高摄取。对于骨、脂肪和肌肉等区域，仅惩罚预测 PET 的正向高值；对于子宫或盆腔区域可设置较弱极值惩罚；膀胱和肠道则作为生理高摄取区域豁免。")
    add_formula(doc, "L_organ = Σ_c η_c · mean(ReLU(ŷ_0) · M_cold^c)")
    add_para(doc, "假热点抑制损失在非病灶且非生理高摄取区域内寻找预测 PET 的高分位异常响应，并对其幅值进行惩罚，防止模型在骨盆背景或肌肉脂肪区域生成孤立热点。")
    add_formula(doc, "L_false = mean( |ŷ_0(p)| · 1{ŷ_0(p)>Q_q(ŷ_0⊙M_cold)} · M_cold(p) )")
    add_para(doc, "PatchNCE 损失通过相同空间位置的 CT patch 与合成 PET patch 构造正样本，不同位置构造负样本，使生成 PET 与 CT 解剖结构保持空间绑定。")
    add_formula(doc, "L_NCE = - log exp(sim(z_i^CT,z_i^PET)/τ_n) / Σ_j exp(sim(z_i^CT,z_j^PET)/τ_n)")
    add_para(doc, "异方差负对数似然损失利用模型输出的 log-variance 对误差进行自适应加权，使噪声较高或映射不确定的位置具有合理的置信表达。")
    add_formula(doc, "L_NLL = 0.5 · exp(-s) · (ŷ_0-x_0)^2 + 0.5 · s,  s=clip(logσ_y^2)")
    add_para(doc, "最终总体损失为：")
    add_formula(doc, "L_total = L_base + λ1L_topk + λ2L_freq + λ3L_NCE + λ4L_SUV + λ5L_false + λ6L_organ + λ7L_hot + λ8L_NLL")
    add_para(doc, "在优选实施方式中，各损失项还乘以阶段门控 g_i(τ)。例如 Top-K、ROI-SUV 和假热点抑制在晚期去噪阶段更强，热点先验监督在中期阶段更强，器官一致性在早中期阶段更强。这样可以使模型先确定全局代谢布局，再逐步恢复小病灶细节和 SUV 峰值。")

    add_heading(doc, "7. 方法流程", 1)
    add_para(doc, "本发明的方法流程可概括为以下步骤。")
    steps = [
        "步骤 S1：获取患者 CT 图像，并可选获取与其配准的 PET 图像、病灶掩膜、器官掩膜、DICOM 元数据和临床 SUV 标定信息。",
        "步骤 S2：对 CT 图像执行 HU 裁剪、尺寸归一化和强度归一化；对 PET 图像执行 SUV 计算和归一化；构建 μ_map、器官 one-hot 掩膜、器官距离图和病例元数据。",
        "步骤 S3：利用 CT 多尺度编码器提取四个尺度的解剖特征，利用 Gabor 先验模块提取高频响应和能量图，利用热点先验模块预测候选热点图，利用器官先验模块提取器官结构特征。",
        "步骤 S4：在训练阶段，按尺度自适应布朗桥扩散公式生成 x_t；在推理阶段，从接近 CT 端点的桥状态开始反向采样。",
        "步骤 S5：将 x_t、源 CT、自条件预测和条件包输入异方差 U-Net，通过 Zero-Conv 注入模块在解码跳跃连接处融合多源先验，得到预测 PET 及不确定性图。",
        "步骤 S6：根据扩散时间步计算阶段化损失，包括基础重建、小病灶聚焦、频域恢复、SUV 定量、器官一致性、假热点抑制和异方差似然等项，并反向更新模型参数。",
        "步骤 S7：推理时采用 DDIM 风格少步采样，逐步从 CT 端桥状态移动到 PET 端预测，输出合成 PET、热点图和置信度图。",
    ]
    for s in steps:
        add_bullet(doc, s)

    add_heading(doc, "8. 系统组成", 1)
    add_table(doc, [
        ["模块", "功能", "可替代实施方式"],
        ["数据预处理模块", "完成 CT/PET 配准质控、HU/SUV 归一化、μ_map 与元数据构建", "可替换为三维体数据预处理或中心化标准化流程"],
        ["CT 编码模块", "提取多尺度解剖特征", "可替换为 ResNet、Swin Transformer 或 3D 编码器"],
        ["Gabor 先验模块", "提取高频纹理和边缘能量", "可替换为小波、LoG、Sobel 或可学习频域滤波器"],
        ["器官先验模块", "提供器官类型、距离和代谢冷/热区域约束", "可由 TotalSegmentator 或人工器官标签产生"],
        ["热点先验模块", "预测病灶候选热区", "可由轻量 U-Net、注意力网络或弱监督热图网络实现"],
        ["尺度自适应噪声模块", "按频带和空间能量分配布朗桥噪声", "可采用不同频带数量或非线性噪声日程"],
        ["Zero-Conv 注入模块", "轻量注入多源先验到 U-Net 跳跃连接", "可替换为门控卷积、FiLM 或低秩适配器"],
        ["去噪生成模块", "预测无噪 PET 与不确定性", "可采用 2D、2.5D 或 3D U-Net 主干"],
        ["损失重分配模块", "实现阶段化代谢保真和小病灶聚焦训练", "可按任务裁剪 SUV、器官、分割一致性等损失"],
    ], [1.15, 3.0, 2.15])

    add_heading(doc, "9. 有益效果", 1)
    add_para(doc, "与传统 CT-to-PET 合成方法相比，本发明至少具有以下有益效果。第一，通过布朗桥扩散将源 CT 明确作为生成路径端点，降低普通无条件或弱条件扩散模型在跨模态生成中的结构漂移风险。第二，通过拉普拉斯金字塔与 Gabor 能量调制实现频带差异化加噪，可减少小病灶边缘和高摄取热点在扩散过程中的过度破坏。第三，通过 Zero-Conv 轻量注入多源医学先验，使模型能够在浅层关注高频边缘，在深层关注器官布局和代谢合理性，并保持较低参数增量。第四，通过 ROI-SUV、Top-K 病灶、频域、器官一致性和假热点抑制等损失项，使生成目标从单纯视觉相似转向临床代谢保真。第五，异方差不确定性输出和 Monte Carlo 采样可为后续人工阅片或下游分析提供置信度参考。")
    add_para(doc, "本发明尤其适用于 PET 数据采集受限、跨中心数据不完整、需要低成本代谢影像估计或需要为分割/检出模型提供代谢先验的场景。需要说明的是，合成 PET 不能替代正式 PET/CT 检查的临床诊断结论，但可作为科研分析、辅助提示和模型预训练的数据补全手段。")

    add_heading(doc, "10. 具体实施方式", 1)
    add_heading(doc, "10.1 实施例一：二维切片 CT-to-PET 合成", 2)
    add_para(doc, "在本实施例中，输入为二维 CT 切片、对应 PET 切片和标签图。图像通过文件名进行配对，统一缩放至 256×256，CT 与 PET 均映射到 [-1,1]。训练配置可采用 batch size 为 4，梯度累积为 2，学习率为 1e-4，权重衰减为 0.01，EMA 衰减为 0.999，扩散时间步为 1000，推理采样步数为 20。启用 Gabor 先验、热点先验、尺度自适应桥噪声、Zero-Conv 条件注入、条件 dropout、自条件机制和异方差输出。")
    add_para(doc, "该实施例适合在已有 PNG 数据集上快速验证模型是否能够完成前向训练、反向采样和可视化输出。由于 PNG 数据通常缺少 SUV 物理标定和器官掩膜，本实施例可先启用 Top-K 病灶损失、焦点频域损失、PatchNCE 损失、热点先验监督和异方差似然损失，暂不启用 ROI-SUV、器官一致性和假热点抑制损失。")

    add_heading(doc, "10.2 实施例二：DICOM/NPZ 临床定量增强合成", 2)
    add_para(doc, "在本实施例中，输入为 DICOM PET/CT 数据。预处理模块从 DICOM 中读取 HU、PET 活度、患者体重、注射剂量、采集时间、半衰期、层厚、层位置和年龄等字段，计算 SUV 图并保存为 NPZ 缓存；同时利用 CT 体数据进行器官分割，生成六类器官掩膜和距离图。该实施例可启用完整 SLMF-BBDM 配置，包括 organ prior、ROI-SUV、false hotspot、organ consistency、metadata FiLM 和可选 segmenter consistency。")
    add_para(doc, "在训练时，模型不仅学习 PET 外观，还受到 SUVmax、SUVmean 和 TBR 的约束；在器官冷区中，模型被惩罚生成不合理高摄取；在膀胱和肠道中，则允许生理性高摄取存在。这种分区约束更符合盆腔 PET/CT 的临床成像规律。")

    add_heading(doc, "10.3 实施例三：不确定性与置信图输出", 2)
    add_para(doc, "在本实施例中，模型启用异方差输出头，并在推理时执行多次独立采样。多次采样的方差作为认知不确定性，模型预测的 log-variance 经指数变换作为偶然不确定性，两者相加得到总不确定性。随后按最大值归一化得到置信度图。该置信图可用于提示医生或研究人员关注低置信区域，也可用于下游分割模型对合成 PET 特征进行加权。")

    add_heading(doc, "11. 附图说明", 1)
    add_para(doc, "图1 为 SLMF-BBDM 总体网络结构示意图，展示 CT 编码、多源先验、Zero-Conv 注入、布朗桥扩散和 U-Net 去噪关系。")
    add_para(doc, "图2 为训练与推理流程图，展示真实 PET、源 CT、前向桥式加噪、条件构建、反向采样和合成 PET 输出。")
    add_para(doc, "图3 为核心创新点示意图，展示频率保护、医学先验轻量注入和代谢保真约束之间的关系。")
    add_para(doc, "图4 为 PET/CT 数据配准与质控流程示意图，展示模型训练前的数据一致性检查。")

    add_heading(doc, "12. 权利要求书建议", 1)
    claims = [
        "一种面向子宫内膜癌 CT 图像的 PET 合成方法，其特征在于，包括：获取 CT 图像；构建 CT 多尺度特征、Gabor 高频能量图和至少一种医学先验；基于源 CT 与目标 PET 构造布朗桥扩散状态；通过条件去噪网络预测 PET 图像；其中所述布朗桥扩散状态采用尺度自适应噪声分配，高频噪声根据 Gabor 能量图进行空间调制。",
        "根据权利要求1所述的方法，其中所述尺度自适应噪声分配包括将图像分解为低频、中频和高频分量，对不同频带设置不同噪声倍率，并在高频分量中降低 Gabor 高能量区域的噪声强度。",
        "根据权利要求1所述的方法，其中所述医学先验包括器官掩膜、器官距离图、CT 衰减映射、热点候选图、语义 token 或病例元数据中的一种或多种。",
        "根据权利要求1所述的方法，其中所述条件去噪网络为 U-Net，并通过逐尺度零初始化 1×1 卷积将多源先验注入所述 U-Net 的跳跃连接。",
        "根据权利要求4所述的方法，其中所述逐尺度零初始化 1×1 卷积的输出乘以与扩散时间相关的注入系数，使不同尺度先验在不同去噪阶段具有不同权重。",
        "根据权利要求1所述的方法，其中训练损失包括基础重建损失、Top-K 病灶聚焦损失、焦点频域损失、PatchNCE 空间绑定损失、ROI-SUV 代谢定量损失、假热点抑制损失、器官一致性损失、热点先验监督损失和异方差负对数似然损失中的至少两种。",
        "根据权利要求6所述的方法，其中 ROI-SUV 代谢定量损失在 SUV 物理空间中约束病灶区域的 SUVmax、SUVmean 和肿瘤背景比。",
        "根据权利要求6所述的方法，其中假热点抑制损失根据器官类型构造代谢冷区掩膜，并惩罚冷区内高于预设分位阈值的预测 PET 响应。",
        "一种 PET 合成系统，包括数据预处理模块、CT 编码模块、频率先验模块、医学先验模块、尺度自适应布朗桥噪声模块、条件注入模块、去噪生成模块和损失计算模块，所述系统用于执行权利要求1至8任一项所述的方法。",
        "一种电子设备，包括处理器和存储器，所述存储器中存储有计算机程序，所述计算机程序被处理器执行时实现权利要求1至8任一项所述的方法。",
        "一种计算机可读存储介质，其上存储有计算机程序，所述计算机程序被执行时实现权利要求1至8任一项所述的方法。",
    ]
    for i, claim in enumerate(claims, start=1):
        add_para(doc, f"{i}. {claim}", first_line=False)

    add_heading(doc, "13. 可保护范围与替代方案", 1)
    add_para(doc, "本发明的保护范围不应限于具体网络通道数、图像尺寸、损失权重或采样步数。CT 编码器可替换为其他卷积网络、Transformer 编码器或 3D 编码器；Gabor 先验可替换为小波、拉普拉斯、方向梯度或其他可学习频域滤波器；器官先验可由自动分割、人工标注或弱监督器官定位模块获得；热点先验可由二分类热图网络、多尺度注意力网络或分割网络产生；扩散主干可采用 DDIM、DPM-Solver 或其他少步采样策略。")
    add_para(doc, "在三维体数据实施方式中，可将二维卷积、二维 Gabor 和二维拉普拉斯金字塔扩展为三维操作，以更好地保持跨层连续性；在多模态实施方式中，还可接入 MRI、临床文本、病理分型或血清学指标作为语义 token 或元数据条件。只要其核心思想仍为基于源 CT 与目标 PET 的桥式扩散、面向小病灶的频率保护噪声、以及多源医学先验的轻量条件注入与代谢保真约束，均应视为本发明构思的等同变形。")

    add_heading(doc, "14. 验证方案与实施注意事项", 1)
    add_para(doc, "为了证明本发明的技术效果，可在不改变发明构思的前提下设计分阶段验证方案。第一阶段验证数据链路和生成主干，即检查 CT、PET、标签、器官图和元数据是否完成患者级划分，确认训练集、验证集和测试集之间不存在同一患者泄漏；同时检查 CT/PET 轮廓 Dice、HD95、切片位置差和图像尺寸一致性，避免模型把配准误差学习为模态差异。第二阶段验证基础 CT-to-PET 生成能力，即在固定训练集上比较普通 U-Net、GAN、标准 DDPM、普通 BBDM 与本发明 SLMF-BBDM 的 MAE、PSNR、SSIM、病灶区域 MAE 和热点定位误差。第三阶段验证临床代谢保真能力，即在具有 SUV 标定的病例上比较 SUVmax 误差、SUVmean 误差、TBR 误差和病灶 Top-K 像素误差。第四阶段验证安全性和可解释性，即观察不确定性图是否在配准差、器官边界、低剂量噪声或异常高摄取区域给出较低置信度。")
    add_para(doc, "消融实验可围绕两个创新点展开。对于创新点一，可依次移除 Gabor 能量调制、移除拉普拉斯多频带噪声、将高频噪声倍率改为与低频相同、将布朗桥扩散替换为普通 DDPM，以验证频率保护和源 CT 端点约束的贡献。对于创新点二，可依次移除 Zero-Conv 条件注入、移除热点先验、移除器官先验、移除 ROI-SUV 损失、移除假热点抑制损失、移除异方差不确定性头，以验证多源医学先验和阶段化误差重分配的贡献。若需要更严谨的专利实施例，也可将完整模型与每个消融模型在相同随机种子、相同训练轮数和相同采样步数下重复训练多次，报告均值、标准差和统计显著性。")
    add_para(doc, "在模型训练过程中，应避免把未配准的 PET/CT 切片、无 SUV 可靠标定的 PET 图像和错误器官掩膜直接用于完整临床约束训练。对于缺少 SUV 字段的样本，可继续用于视觉生成和频域约束，但 ROI-SUV 损失应自动跳过；对于缺少器官掩膜的样本，可用零器官通道作为占位，以保证模型接口一致；对于热点标签不完整的样本，可由 PET 高摄取分位区域构建软监督，但不宜把该区域等同于病理确诊病灶。")
    add_para(doc, "在推理部署时，系统可输出合成 PET、热点候选图、像素级不确定性和置信度图。为了降低误用风险，合成 PET 应标注为算法生成影像，不能替代真实 PET/CT 的诊断证据。对于置信度低、病灶周围结构复杂或存在明显配准异常的病例，系统应提示需要人工复核或建议结合真实 PET、MRI、病理和临床资料综合判断。若用于科研数据补全，可记录模型版本、训练数据范围、预处理窗口、采样步数和不确定性阈值，以便后续复现实验和审查数据来源。")
    add_para(doc, "本发明的工程实现还可设置若干质量控制阈值。例如，当 CT 与 PET 大轮廓重叠度低于预设阈值、切片 z 位置差超过预设毫米数、SUV 标定字段不完整、器官掩膜覆盖率异常或模型输出不确定性大面积升高时，系统可拒绝输出临床可用图像，只保留研究提示结果。该机制并非发明核心步骤的限制，而是为了在真实医院数据环境中提高系统可靠性。")

    add_heading(doc, "15. 产业应用场景", 1)
    add_para(doc, "本发明可应用于多种医疗和科研场景。其一，在 PET 检查资源不足或患者暂不适合 PET 扫描时，可基于 CT 生成代谢影像提示图，为研究人员提供辅助参考。其二，在多中心回顾性研究中，不同中心可能只保存 CT 或 PET 缺失严重，本发明可作为跨模态数据补全工具，帮助构建更完整的影像组学或深度学习研究队列。其三，在子宫内膜癌术前评估中，合成 PET 可与 CT 解剖图共同输入后续分割或检出模型，提高模型对高摄取疑似病灶区域的关注。其四，在教学和算法验证中，本发明可生成带有置信度提示的 PET 风格图像，辅助展示 CT 结构与代谢分布之间的对应关系。")
    add_para(doc, "从产品形态看，本发明可部署为医院内网服务器、影像工作站插件、科研平台批处理工具或云端推理服务。系统可读取 DICOM、NIfTI、PNG 或 NPZ 格式数据，经过预处理后输出合成 PET DICOM、普通图像、数值张量、热点掩膜和报告表格。若与 PACS/RIS 系统连接，可将合成结果作为二级派生图层显示，并在界面中保留“算法生成”“非诊断替代”的水印和元数据。")
    add_para(doc, "在数据管理方面，本发明可为每次推理建立完整记录，包括输入影像编号、预处理参数、器官先验来源、模型权重版本、采样步数、随机种子、输出文件哈希值和置信度统计。上述记录既有助于科研复现，也有助于后续专利实施中的质量追踪。若在不同医院或不同扫描协议下部署，可通过少量本地数据进行微调，也可仅更新归一化窗口、SUV 上限、器官映射规则和不确定性阈值，以减少跨中心域偏移。")
    add_para(doc, "从权利布局角度看，可围绕方法、系统、装置、介质和模型训练方法分别撰写权利要求。独立权利要求宜覆盖“布朗桥 CT-to-PET 合成、频带自适应噪声、Gabor 或等价高频能量调制、医学先验轻量注入、代谢保真阶段化损失”这几个共同技术特征；从属权利要求可进一步限定器官类别、热点先验训练方式、ROI-SUV 损失、假热点抑制、不确定性图输出和三维扩展。这样既能保护当前工程实现，也能覆盖后续模型升级。")
    add_para(doc, "预期技术效果方面，本发明并非单纯追求生成图像外观相似，而是把微小病灶可见性、代谢峰值保持、器官生理摄取区分和输出置信表达同时纳入模型设计。通过该综合约束，系统有望减少小病灶被背景平均化、正常膀胱或肠道摄取被误判、冷区出现孤立假热点以及低可信结果无提示输出等问题，并提高后续人工复核和科研分析的针对性、稳定性与可追溯性。")

    add_heading(doc, "16. 结论", 1)
    add_para(doc, "本发明围绕子宫内膜癌 CT-to-PET 合成中的小病灶易平滑、代谢热点易错位、器官生理摄取易混淆和临床 SUV 指标难保持等问题，提出了一种 SLMF-BBDM 技术方案。该方案以布朗桥扩散为跨模态生成骨架，以 Gabor 能量和拉普拉斯频带分解实现尺度自适应噪声保护，以 Zero-Conv 实现多源医学先验轻量注入，并以阶段化损失函数实现病灶、频域、器官和 SUV 四个层面的误差重分配。由此，模型能够在理论结构和工程实现上兼顾生成质量、医学可解释性和临床定量约束，适合作为专利申请中的方法、系统、设备和存储介质方案进行布局。")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build_doc())
