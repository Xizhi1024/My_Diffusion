from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
import numpy as np
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle


ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = ROOT.parent / "data"
OUT_DIR = ROOT / "output" / "midterm_report"
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "工作3中期答辩进展报告_SLMF-BBDM_Word版.docx"

NAVY = "102A43"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "0F766E"
CYAN = "0891B2"
ORANGE = "D97706"
PURPLE = "6D5BD0"
RED = "B91C1C"
GRAY = "5F6C7B"
LIGHT_GRAY = "F2F4F7"
LINE = "D7E0EA"


def hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = hex_to_rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000",
                  size: float = 9.5, align: int | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_body_para(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        set_run_font(run)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, color=BLUE, bold=True)
    elif level == 2:
        set_run_font(run, size=13, color=BLUE, bold=True)
    else:
        set_run_font(run, size=12, color=DARK_BLUE, bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=GRAY, italic=True)


def count_files(folder: Path, sub: str) -> int:
    path = folder / sub
    return len([p for p in path.glob("*") if p.is_file()]) if path.exists() else 0


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def load_yaml(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {}


def draw_box(ax, x, y, w, h, title, subtitle="", fc="#FFFFFF", ec="#2E74B5",
             lw=1.2, dashed=False, title_size=8.5, sub_size=7.2) -> None:
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.018,rounding_size=0.025",
        linewidth=lw,
        edgecolor=ec,
        facecolor=fc,
        linestyle="--" if dashed else "-",
    )
    ax.add_patch(patch)
    title_lines = title.count("\n") + 1
    title_y = y + h - 0.018
    ax.text(x + w / 2, title_y, title, ha="center", va="top",
            fontsize=title_size, color="#102A43", fontweight="bold", linespacing=0.92)
    if subtitle:
        ax.text(x + w / 2, title_y - 0.025 * title_lines, subtitle, ha="center", va="top",
                fontsize=sub_size, color="#455A6B", linespacing=1.12)


def arrow(ax, start, end, color="#2E74B5", lw=1.2, dashed=False, rad=0.0) -> None:
    arr = FancyArrowPatch(
        start, end,
        arrowstyle="-|>",
        mutation_scale=10,
        linewidth=lw,
        color=color,
        linestyle="--" if dashed else "-",
        shrinkA=3,
        shrinkB=3,
        connectionstyle=f"arc3,rad={rad}",
    )
    ax.add_patch(arr)


def make_pipeline_figure(out_path: Path) -> None:
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 8,
        "axes.linewidth": 0.8,
        "figure.dpi": 180,
    })
    fig, ax = plt.subplots(figsize=(12.8, 6.8))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.text(0.02, 0.97, "SLMF-BBDM training and inference graph (PNG256 implementation)",
            fontsize=11.5, fontweight="bold", color="#102A43", va="top")
    ax.text(0.02, 0.935, "Solid arrows: active path in slmf_png256.yaml. Dashed boxes: implemented but disabled in the current PNG run.",
            fontsize=8, color="#5F6C7B", va="top")

    # Vertical zones.
    zones = [
        (0.02, 0.10, 0.18, 0.78, "Data interface", "#F7FAFC"),
        (0.225, 0.10, 0.45, 0.78, "Bridge diffusion core", "#FBFCFE"),
        (0.70, 0.10, 0.28, 0.78, "Objectives / output", "#F7FAFC"),
    ]
    for x, y, w, h, label, fc in zones:
        ax.add_patch(Rectangle((x, y), w, h, facecolor=fc, edgecolor="#D7E0EA", linewidth=1.0))
        ax.text(x + 0.012, y + h - 0.035, label, fontsize=9, fontweight="bold", color="#102A43")

    # Data.
    draw_box(ax, 0.045, 0.755, 0.072, 0.07, "CT PNG", "512x512", "#EAF2FF", "#2E74B5", title_size=7.4, sub_size=6.5)
    draw_box(ax, 0.126, 0.755, 0.072, 0.07, "PET PNG", "192x192", "#E8F7F2", "#0F766E", title_size=7.4, sub_size=6.5)
    draw_box(ax, 0.045, 0.635, 0.153, 0.07, "Lesion label", "paired filename", "#FFF7E8", "#D97706", title_size=7.6, sub_size=6.4)
    draw_box(ax, 0.045, 0.49, 0.153, 0.09, "Online preprocessing", "normalize + resize\n256x256 tensor", "#FFFFFF", "#455A6B", title_size=7.4, sub_size=6.2)
    draw_box(ax, 0.045, 0.34, 0.153, 0.085, "Model batch", "x_s: CT; x_0: PET; M: mask", "#E8F7F2", "#0F766E", title_size=7.4, sub_size=6.2)
    draw_box(ax, 0.045, 0.18, 0.153, 0.075, "Metadata", "sample / patient / slice\nnative PET size", "#FFF7E8", "#D97706", title_size=7.4, sub_size=6.1)
    arrow(ax, (0.077, 0.75), (0.077, 0.70), "#455A6B")
    arrow(ax, (0.162, 0.75), (0.162, 0.70), "#455A6B")
    arrow(ax, (0.12, 0.63), (0.12, 0.58), "#455A6B")
    arrow(ax, (0.12, 0.49), (0.12, 0.425), "#455A6B")
    arrow(ax, (0.12, 0.34), (0.12, 0.255), "#D97706", dashed=True)

    # Core.
    draw_box(ax, 0.245, 0.76, 0.115, 0.07, "Target PET", "x_0", "#E8F7F2", "#0F766E", title_size=7.6, sub_size=6.6)
    draw_box(ax, 0.245, 0.64, 0.115, 0.07, "Source CT", "x_s", "#EAF2FF", "#2E74B5", title_size=7.6, sub_size=6.6)
    draw_box(ax, 0.405, 0.69, 0.145, 0.09, "Scale-adaptive\nBBDM bridge", "band-wise sigma\n+ bridge mixing", "#FFF7E8", "#A16207", title_size=7.2, sub_size=6.0)
    draw_box(ax, 0.585, 0.70, 0.085, 0.07, "Noisy state", "x_t", "#E8F7F2", "#0F766E", title_size=7.2, sub_size=6.3)
    draw_box(ax, 0.405, 0.505, 0.195, 0.11, "BBDM U-Net denoiser", "[x_t, x_s, self-cond]\nZero-conv skip injection", "#E8F7F2", "#0F766E", title_size=7.5, sub_size=6.2)
    draw_box(ax, 0.620, 0.50, 0.055, 0.075, "Output", "x0_hat\nlogvar", "#E8F7F2", "#0F766E", title_size=6.8, sub_size=5.8)

    arrow(ax, (0.365, 0.795), (0.405, 0.74), "#0F766E")
    arrow(ax, (0.365, 0.675), (0.405, 0.72), "#2E74B5")
    arrow(ax, (0.550, 0.735), (0.585, 0.735), "#A16207")
    arrow(ax, (0.585, 0.70), (0.51, 0.615), "#0F766E")
    arrow(ax, (0.365, 0.64), (0.405, 0.56), "#2E74B5", rad=-0.05)
    arrow(ax, (0.605, 0.555), (0.615, 0.54), "#0F766E")
    arrow(ax, (0.195, 0.382), (0.245, 0.675), "#2E74B5", rad=-0.12)
    arrow(ax, (0.195, 0.382), (0.245, 0.795), "#0F766E", rad=-0.16)

    # Priors and condition path.
    draw_box(ax, 0.245, 0.48, 0.115, 0.075, "CT encoder", "4-scale features", "#EAF2FF", "#2E74B5", title_size=7.3, sub_size=6.2)
    draw_box(ax, 0.245, 0.34, 0.115, 0.075, "Gabor prior", "32 filters", "#EAF2FF", "#2E74B5", title_size=7.3, sub_size=6.2)
    draw_box(ax, 0.405, 0.34, 0.115, 0.075, "Hotspot prior", "CT-to-lesion", "#EAF2FF", "#2E74B5", title_size=7.3, sub_size=6.2)
    draw_box(ax, 0.555, 0.31, 0.12, 0.095, "Condition bundle", "CT feats + Gabor\n+ hotspot map", "#FFFFFF", "#455A6B", title_size=7.2, sub_size=6.0)
    draw_box(ax, 0.555, 0.18, 0.12, 0.08, "Zero-conv adapter", "beta(t) skip\ninjections", "#FFFFFF", "#455A6B", title_size=7.0, sub_size=5.9)
    draw_box(ax, 0.245, 0.21, 0.115, 0.055, "Organ prior", "disabled", "#FFFFFF", "#98A2B3", dashed=True, title_size=7.0, sub_size=5.8)
    draw_box(ax, 0.405, 0.21, 0.115, 0.055, "Semantic tokens", "disabled", "#FFFFFF", "#98A2B3", dashed=True, title_size=7.0, sub_size=5.8)

    arrow(ax, (0.305, 0.64), (0.305, 0.555), "#2E74B5")
    arrow(ax, (0.305, 0.48), (0.305, 0.415), "#2E74B5")
    arrow(ax, (0.365, 0.377), (0.405, 0.377), "#2E74B5")
    arrow(ax, (0.305, 0.34), (0.555, 0.36), "#2E74B5", rad=0.05)
    arrow(ax, (0.465, 0.34), (0.555, 0.35), "#2E74B5")
    arrow(ax, (0.615, 0.31), (0.615, 0.26), "#455A6B")
    arrow(ax, (0.615, 0.26), (0.515, 0.51), "#455A6B", rad=0.15)
    arrow(ax, (0.300, 0.415), (0.405, 0.71), "#A16207", rad=-0.10)
    arrow(ax, (0.305, 0.265), (0.555, 0.33), "#98A2B3", dashed=True, rad=0.12)
    arrow(ax, (0.465, 0.265), (0.555, 0.335), "#98A2B3", dashed=True, rad=0.05)

    # Objectives.
    draw_box(ax, 0.72, 0.68, 0.22, 0.13, "Training objectives", "MSE/L1/grad + Min-SNR\nNLL + lesion/frequency/PatchNCE\nhotspot supervision", "#FFF7E8", "#A16207", title_size=7.4, sub_size=6.3)
    draw_box(ax, 0.72, 0.50, 0.22, 0.07, "Disabled in PNG256", "ROI-SUV, organ, false hotspot,\nsegmenter, metadata FiLM", "#FFFFFF", "#98A2B3", dashed=True, title_size=7.2, sub_size=5.9)
    draw_box(ax, 0.72, 0.34, 0.22, 0.065, "Inference input", "CT PNG -> 256x256", "#EAF2FF", "#2E74B5", title_size=7.4, sub_size=6.2)
    draw_box(ax, 0.72, 0.235, 0.22, 0.065, "DDIM reverse bridge", "x_T -> ... -> x0_hat | CT", "#E8F7F2", "#0F766E", title_size=7.4, sub_size=6.2)
    draw_box(ax, 0.72, 0.135, 0.22, 0.065, "Predicted PET PNG", "native PET size", "#E8F7F2", "#0F766E", title_size=7.4, sub_size=6.2)
    arrow(ax, (0.675, 0.54), (0.72, 0.72), "#A16207")
    arrow(ax, (0.83, 0.34), (0.83, 0.30), "#455A6B")
    arrow(ax, (0.83, 0.235), (0.83, 0.20), "#455A6B")
    arrow(ax, (0.83, 0.135), (0.83, 0.12), "#0F766E")

    # Legend.
    ax.add_patch(Rectangle((0.04, 0.065), 0.022, 0.014, facecolor="#E8F7F2", edgecolor="#0F766E"))
    ax.text(0.066, 0.066, "active tensors / trainable modules", fontsize=7, va="bottom", color="#455A6B")
    ax.add_patch(Rectangle((0.235, 0.065), 0.022, 0.014, facecolor="#FFF7E8", edgecolor="#A16207"))
    ax.text(0.261, 0.066, "diffusion / objectives", fontsize=7, va="bottom", color="#455A6B")
    ax.add_patch(Rectangle((0.385, 0.065), 0.022, 0.014, facecolor="#FFFFFF", edgecolor="#98A2B3", linestyle="--"))
    ax.text(0.411, 0.066, "implemented but disabled", fontsize=7, va="bottom", color="#455A6B")
    ax.text(0.98, 0.066, "Source: configs/experiments/slmf_png256.yaml + src/model/*.py",
            fontsize=7, va="bottom", ha="right", color="#5F6C7B")
    fig.savefig(out_path, dpi=320, bbox_inches="tight")
    plt.close(fig)


def make_data_qc_figure(out_path: Path, train: dict, val: dict) -> None:
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8, "axes.linewidth": 0.8})
    fig, axes = plt.subplots(2, 2, figsize=(8.2, 5.6))
    fig.suptitle("Data split and PET/CT registration quality control", x=0.02, ha="left",
                 fontsize=12, fontweight="bold", color="#102A43")

    # A: counts.
    ax = axes[0, 0]
    labels = ["Train", "Val"]
    slices = [train.get("images", 953), val.get("images", 238)]
    patients = [train.get("patients", 124), val.get("patients", 31)]
    x = np.arange(2)
    ax.bar(x - 0.17, slices, width=0.34, color="#2E74B5", label="slices")
    ax.bar(x + 0.17, patients, width=0.34, color="#0F766E", label="patients")
    ax.set_xticks(x, labels)
    ax.set_ylabel("Count")
    ax.set_title("A. Dataset scale", loc="left", fontweight="bold")
    ax.legend(frameon=False, ncol=2, fontsize=7)
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.6)

    # B: body Dice.
    ax = axes[0, 1]
    before = [train["before_body_dice"]["mean"], val["before_body_dice"]["mean"]]
    after = [train["after_body_dice"]["mean"], val["after_body_dice"]["mean"]]
    ax.plot(labels, before, "o--", color="#98A2B3", label="before")
    ax.plot(labels, after, "o-", color="#0F766E", label="after")
    for i, y in enumerate(after):
        ax.text(i, y + 0.012, f"{y:.3f}", ha="center", fontsize=7, color="#0F766E")
    ax.set_ylim(0.65, 1.0)
    ax.set_ylabel("Body contour Dice")
    ax.set_title("B. Field-of-view alignment improves", loc="left", fontweight="bold")
    ax.legend(frameon=False, fontsize=7)
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.6)

    # C: HD95.
    ax = axes[1, 0]
    before_hd = [train["before_contour_hd95_pixels"]["mean"], val["before_contour_hd95_pixels"]["mean"]]
    after_hd = [train["after_contour_hd95_pixels"]["mean"], val["after_contour_hd95_pixels"]["mean"]]
    ax.plot(labels, before_hd, "o--", color="#B91C1C", label="before")
    ax.plot(labels, after_hd, "o-", color="#2E74B5", label="after")
    for i, y in enumerate(after_hd):
        ax.text(i, y + 1.1, f"{y:.2f}", ha="center", fontsize=7, color="#2E74B5")
    ax.set_ylabel("HD95 (pixels)")
    ax.set_title("C. Boundary distance drops", loc="left", fontweight="bold")
    ax.legend(frameon=False, fontsize=7)
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.6)

    # D: pass rates.
    ax = axes[1, 1]
    pass_rate = [train["slice_pass_rate"] * 100, val["slice_pass_rate"] * 100]
    all_patient = [
        train["patients_with_all_slices_passing"] / train["patients"] * 100,
        val["patients_with_all_slices_passing"] / val["patients"] * 100,
    ]
    ax.bar(x - 0.17, pass_rate, width=0.34, color="#0F766E", label="slice pass rate")
    ax.bar(x + 0.17, all_patient, width=0.34, color="#6D5BD0", label="all-slice patient pass")
    for i, y in enumerate(pass_rate):
        ax.text(i - 0.17, y + 1.0, f"{y:.1f}%", ha="center", fontsize=7)
    ax.set_xticks(x, labels)
    ax.set_ylim(0, 108)
    ax.set_ylabel("%")
    ax.set_title("D. QC pass rate", loc="left", fontweight="bold")
    ax.legend(frameon=False, fontsize=7)
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.6)

    for ax in axes.flat:
        ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    fig.savefig(out_path, dpi=320, bbox_inches="tight")
    plt.close(fig)


def make_status_figure(out_path: Path) -> None:
    components = [
        ("Data manifest / split", "Done", "Patient-level split and PNG pairing"),
        ("PET registration QC", "Done", "Train and val field-of-view checks"),
        ("PNGSliceDataset", "Done", "256x256 online tensors + metadata"),
        ("SLMF-BBDM core", "Done", "pred_x0 BBDM U-Net, sampling"),
        ("Gabor + hotspot priors", "Enabled", "Current PNG256 config"),
        ("Scale-adaptive bridge", "Enabled", "Band-wise sigma + Gabor modulation"),
        ("Zero-conv adapter", "Enabled", "Time-varying skip injection"),
        ("Clinical SUV / organ losses", "Ready/off", "Implemented, needs SUV/DICOM cache"),
        ("Semantic / segmenter branch", "Ready/off", "Code path exists, not enabled"),
        ("Formal checkpoint + eval", "Next", "No checkpoint/results found in workdir"),
    ]
    status_color = {"Done": "#0F766E", "Enabled": "#2E74B5", "Ready/off": "#D97706", "Next": "#B91C1C"}
    fig, ax = plt.subplots(figsize=(8.5, 5.0))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.text(0.02, 0.96, "Implementation status matrix", fontsize=12, fontweight="bold",
            color="#102A43", va="top")
    ax.text(0.02, 0.91, "Current state inferred from code, configs, data folders, and test execution.",
            fontsize=8, color="#5F6C7B", va="top")
    y0 = 0.84
    row_h = 0.071
    for i, (name, status, note) in enumerate(components):
        y = y0 - i * row_h
        fill = "#F7FAFC" if i % 2 == 0 else "#FFFFFF"
        ax.add_patch(Rectangle((0.02, y - 0.047), 0.96, 0.058, facecolor=fill, edgecolor="#E5E7EB", linewidth=0.6))
        ax.text(0.045, y - 0.017, name, fontsize=8.2, fontweight="bold", color="#102A43", va="center")
        ax.add_patch(FancyBboxPatch((0.37, y - 0.037), 0.095, 0.038,
                                    boxstyle="round,pad=0.006,rounding_size=0.015",
                                    facecolor=status_color[status], edgecolor=status_color[status]))
        ax.text(0.417, y - 0.018, status, ha="center", va="center", fontsize=7.2,
                color="white", fontweight="bold")
        ax.text(0.49, y - 0.017, note, fontsize=8, color="#455A6B", va="center")
    ax.text(0.98, 0.04, "Engineering validation: pytest tests/test_smoke.py -q => 87 passed",
            fontsize=7.5, color="#0F766E", ha="right")
    fig.savefig(out_path, dpi=320, bbox_inches="tight")
    plt.close(fig)


def make_roadmap_figure(out_path: Path) -> None:
    tasks = [
        ("Freeze PNG256\ntraining run", 0, 2, "#2E74B5"),
        ("Baseline/full\nquantitative eval", 2, 2, "#0F766E"),
        ("Ablation matrix\n6-8 variants", 4, 3, "#6D5BD0"),
        ("DICOM/SUV cache\norgan prior", 6, 3, "#D97706"),
        ("Clinical statistics\nand thesis figures", 8, 4, "#B91C1C"),
    ]
    fig, ax = plt.subplots(figsize=(8.6, 3.2))
    ax.set_title("Next-step roadmap after mid-term review", loc="left",
                 fontsize=12, fontweight="bold", color="#102A43")
    y_pos = np.arange(len(tasks))[::-1]
    for i, (label, start, duration, color) in enumerate(tasks):
        y = y_pos[i]
        ax.barh(y, duration, left=start, height=0.45, color=color, alpha=0.90)
        ax.text(start + duration / 2, y, label, ha="center", va="center",
                fontsize=8, color="white", fontweight="bold")
    ax.set_yticks([])
    ax.set_xlim(0, 12)
    ax.set_xlabel("Weeks after mid-term")
    ax.set_xticks(range(0, 13, 2))
    ax.grid(axis="x", color="#E5E7EB", linewidth=0.7)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.text(0.99, 0.98, "Decision gate: checkpoint + validation metrics before performance claims",
            fontsize=7.2, color="#5F6C7B", ha="right", va="top", transform=ax.transAxes)
    fig.tight_layout(pad=1.1)
    fig.savefig(out_path, dpi=320, bbox_inches="tight")
    plt.close(fig)


def make_innovation_figure(out_path: Path) -> None:
    fig, ax = plt.subplots(figsize=(8.6, 4.2))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.text(0.02, 0.96, "Three method points to refine for the thesis / paper story",
            fontsize=12, fontweight="bold", color="#102A43", va="top")
    cards = [
        ("01", "Scale-adaptive\nBrownian bridge", "CT-to-PET bridge\n+ Laplacian-band noise\n+ Gabor energy protection", "#0F766E", "#E8F7F2"),
        ("02", "Prior-guided\nconditioning", "CT encoder, Gabor,\nhotspot and organ priors\nvia lightweight zero-conv", "#2E74B5", "#EAF2FF"),
        ("03", "Metabolic fidelity\nand uncertainty", "Lesion Top-K, frequency,\nPatchNCE, NLL / MC confidence\nwith SUV extensions ready", "#6D5BD0", "#F1EEFF"),
    ]
    xs = [0.04, 0.36, 0.68]
    for x, (num, title, body, color, fill) in zip(xs, cards):
        ax.add_patch(FancyBboxPatch((x, 0.18), 0.27, 0.66,
                                    boxstyle="round,pad=0.018,rounding_size=0.025",
                                    facecolor=fill, edgecolor=color, linewidth=1.2))
        ax.add_patch(plt.Circle((x + 0.045, 0.77), 0.033, color=color))
        ax.text(x + 0.045, 0.77, num, ha="center", va="center",
                fontsize=10, color="white", fontweight="bold")
        ax.text(x + 0.09, 0.79, title, ha="left", va="top",
                fontsize=8.9, color="#102A43", fontweight="bold", linespacing=0.95)
        ax.text(x + 0.045, 0.56, body, ha="left", va="top",
                fontsize=7.4, color="#455A6B", linespacing=1.45)
        ax.text(x + 0.045, 0.28, "Evidence now: code implemented\nNeeded next: ablation + metrics",
                ha="left", va="top", fontsize=6.5, color=color, linespacing=1.15)
    fig.savefig(out_path, dpi=320, bbox_inches="tight")
    plt.close(fig)


def build_docx() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    train_summary = load_json(DATA_ROOT / "train" / "pet_peizhuan_validation" / "large_contour_validation_summary.json")
    val_summary = load_json(DATA_ROOT / "val" / "pet_peizhuan_validation" / "large_contour_validation_summary.json")
    png_cfg = load_yaml(ROOT / "configs" / "experiments" / "slmf_png256.yaml")
    full_cfg = load_yaml(ROOT / "configs" / "experiments" / "slmf_full.yaml")
    ablations = load_yaml(ROOT / "configs" / "experiments" / "ablations.yaml")

    figures = {
        "pipeline": FIG_DIR / "fig1_slmf_bbdm_pipeline_sci.png",
        "data_qc": FIG_DIR / "fig2_data_registration_qc.png",
        "status": FIG_DIR / "fig3_implementation_status.png",
        "innovation": FIG_DIR / "fig4_innovation_points.png",
        "roadmap": FIG_DIR / "fig5_next_roadmap.png",
    }
    make_pipeline_figure(figures["pipeline"])
    make_data_qc_figure(figures["data_qc"], train_summary, val_summary)
    make_status_figure(figures["status"])
    make_innovation_figure(figures["innovation"])
    make_roadmap_figure(figures["roadmap"])

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = hex_to_rgb(color)
        st.font.bold = True
    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    # Header / footer.
    hp = section.header.paragraphs[0]
    hp.text = ""
    hr = hp.add_run("工作3中期答辩进展报告 | SLMF-BBDM")
    set_run_font(hr, size=9, color=GRAY)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fr = fp.add_run("基于当前工作目录源码、配置、数据质控结果整理 | 2026-06-30")
    set_run_font(fr, size=8.5, color=GRAY)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Cover / masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("研究生中期答辩进展报告")
    set_run_font(run, size=24, color=NAVY, bold=True)
    p = doc.add_paragraph()
    run = p.add_run("工作3：SLMF-BBDM 面向子宫内膜癌 CT-to-PET 合成的小病灶代谢保真扩散模型")
    set_run_font(run, size=13.5, color=GRAY)
    p.paragraph_format.space_after = Pt(14)

    meta = [
        ("研究任务", "利用 CT 图像合成 PET 代谢图像，并尽量保持小病灶、高频边缘、SUV 定量和不确定性信息。"),
        ("代码依据", "当前目录 My_diffusion 下的 src/、configs/、scripts/、tests/、output/ 与相邻 data/ 目录。"),
        ("当前结论", "数据配准与模型工程框架基本完成，PNG256 训练配置已就绪；正式 checkpoint、训练曲线和验证集定量结果尚需跑通。"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.autofit = False
    for i, (k, v) in enumerate(meta):
        set_cell_text(table.cell(i, 0), k, bold=True, color=NAVY, size=10)
        set_cell_text(table.cell(i, 1), v, size=10)
        set_cell_shading(table.cell(i, 0), LIGHT_GRAY)
        for j in range(2):
            set_cell_border(table.cell(i, j))
            table.cell(i, j).width = Inches(1.25 if j == 0 else 5.1)
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("阶段性判断：目前最主要的工作重心应从“继续堆模块”转向“正式训练、消融验证、患者级统计和临床定量证据闭环”。")
    set_run_font(run, size=12, color=RED, bold=True)
    doc.add_page_break()

    add_heading(doc, "1. 一页式进展总览", 1)
    add_body_para(doc, "本阶段工作已经完成了从数据组织、配准质控、模型代码、训练入口、评估脚本到 smoke tests 的工程闭环。当前目录没有发现正式训练 checkpoint 或 evaluation JSON，因此模型性能不能在本报告中提前宣称；但关键代码路径已经通过自动化测试。")
    kpis = [
        ("数据准备", f"train: {train_summary.get('images', 953)} 张 / {train_summary.get('patients', 124)} 人；val: {val_summary.get('images', 238)} 张 / {val_summary.get('patients', 31)} 人"),
        ("配准质控", f"验证集配准后 Dice={val_summary.get('after_body_dice', {}).get('mean', 0):.3f}，HD95={val_summary.get('after_contour_hd95_pixels', {}).get('mean', 0):.2f}px，通过率={val_summary.get('slice_pass_rate', 0)*100:.1f}%"),
        ("模型工程", "SLMF-BBDM、PNGSliceDataset、Gabor prior、Hotspot prior、Scale-adaptive BBDM、Zero-conv adapter 均已实现"),
        ("测试验证", "ECPC 环境下 `pytest tests/test_smoke.py -q`：87 passed in 16.36s"),
    ]
    table = doc.add_table(rows=1 + len(kpis), cols=2)
    for j, txt in enumerate(["维度", "当前状态"]):
        set_cell_text(table.cell(0, j), txt, bold=True, color="FFFFFF", size=10)
        set_cell_shading(table.cell(0, j), NAVY)
        set_cell_border(table.cell(0, j), NAVY)
    for i, row in enumerate(kpis, start=1):
        set_cell_text(table.cell(i, 0), row[0], bold=True, color=NAVY, size=9.8)
        set_cell_text(table.cell(i, 1), row[1], size=9.8)
        set_cell_shading(table.cell(i, 0), LIGHT_GRAY)
        for j in range(2):
            set_cell_border(table.cell(i, j))

    doc.add_picture(str(figures["status"]), width=Inches(6.45))
    add_caption(doc, "图1  当前实现状态矩阵：已完成、已启用、已实现但暂未启用和下一步证据缺口。")
    doc.add_page_break()

    add_heading(doc, "2. 数据准备与配准质控", 1)
    add_body_para(doc, "当前工作目录相邻的 data/ 目录已经整理出直接 PNG 数据接口，配置文件 `configs/experiments/slmf_png256.yaml` 使用 `../data/{train,val}/{images,pet,labels}` 的同名文件配对方式。实际目录中还包含 `ct`、`pet`、`label`、`pet_peizhuan` 四类图像；本报告按训练配置和配准产物共同描述数据侧进展。")
    data_table = doc.add_table(rows=4, cols=5)
    headers = ["Split", "Patients", "CT/Images", "PET", "Labels"]
    for j, h in enumerate(headers):
        set_cell_text(data_table.cell(0, j), h, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(data_table.cell(0, j), NAVY)
        set_cell_border(data_table.cell(0, j), NAVY)
    rows = [
        ("Train", str(train_summary.get("patients", 124)), str(count_files(DATA_ROOT / "train", "ct")), str(count_files(DATA_ROOT / "train", "pet")), str(count_files(DATA_ROOT / "train", "label"))),
        ("Val", str(val_summary.get("patients", 31)), str(count_files(DATA_ROOT / "val", "ct")), str(count_files(DATA_ROOT / "val", "pet")), str(count_files(DATA_ROOT / "val", "label"))),
        ("Total", str(train_summary.get("patients", 124) + val_summary.get("patients", 31)), str(count_files(DATA_ROOT / "train", "ct") + count_files(DATA_ROOT / "val", "ct")), str(count_files(DATA_ROOT / "train", "pet") + count_files(DATA_ROOT / "val", "pet")), str(count_files(DATA_ROOT / "train", "label") + count_files(DATA_ROOT / "val", "label"))),
    ]
    for i, row in enumerate(rows, 1):
        for j, txt in enumerate(row):
            set_cell_text(data_table.cell(i, j), txt, bold=(j == 0), color=NAVY if j == 0 else "000000", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            if j == 0:
                set_cell_shading(data_table.cell(i, j), LIGHT_GRAY)
            set_cell_border(data_table.cell(i, j))

    doc.add_picture(str(figures["data_qc"]), width=Inches(6.45))
    add_caption(doc, "图2  数据规模与 PET/CT 大轮廓配准质控。该质控检验视野级对齐，不等同于 DICOM 毫米级物理配准。")
    add_body_para(doc, "配准质控的阶段性结论：训练集 953 张中 949 张满足全部严格条件，验证集 238 张全部满足；验证集 Dice 从约 0.708 提升到约 0.953，HD95 从约 24.83px 降至约 4.36px。该结果支撑后续 PNG256 二维合成实验，但后续若启用 SUV 和器官先验，仍建议回到 DICOM/NPZ cache 版本补充物理尺度字段。")
    doc.add_page_break()

    add_heading(doc, "3. 当前模型技术路线", 1)
    add_body_para(doc, "SLMF-BBDM 的核心思想是把 CT 和 PET 看作布朗桥的两个端点，在训练时从真实 PET 与源 CT 构造有条件的噪声状态，再由 U-Net 在多源先验约束下预测无噪 PET。当前 PNG256 配置启用了 Gabor 先验、Hotspot 先验、Scale-adaptive bridge、Zero-conv adapter、condition dropout、self-conditioning 与异方差不确定性头。")
    doc.add_picture(str(figures["pipeline"]), width=Inches(6.45))
    add_caption(doc, "图3  SCI 风格训练/推理路径图。实线为 PNG256 当前启用路径，虚线为已实现但当前未启用路径。")

    add_heading(doc, "4. 已实现模块与配置状态", 1)
    add_body_para(doc, "从 `slmf_png256.yaml` 看，当前用于直接 PNG 训练的版本更强调可跑通和小病灶形态保真：启用 Gabor、高频保护、Hotspot prior、PatchNCE、Top-K lesion、focal frequency 与 heteroscedastic NLL；暂时关闭 organ prior、ROI-SUV、false hotspot、semantic prior、metadata FiLM 和 segmenter consistency。")
    enabled_modules = [k for k, v in png_cfg.get("modules", {}).items() if isinstance(v, dict) and v.get("enabled")]
    enabled_losses = [k for k, v in png_cfg.get("losses", {}).items() if isinstance(v, dict) and v.get("enabled")]
    add_bullets(doc, [
        "当前 PNG256 启用模块：" + "、".join(enabled_modules),
        "当前 PNG256 启用损失：" + "、".join(enabled_losses),
        "Full 配置中已准备但当前不一定启用的模块包括 organ prior、ROI-SUV、false hotspot、organ consistency 和 metadata/segmenter 分支。",
    ])
    doc.add_page_break()

    add_heading(doc, "5. 可凝练的三个方法创新点", 1)
    add_body_para(doc, "以下表述是基于当前代码的中期阶段凝练，后续写论文前还需要进一步检索文献，确认“首次提出”或“差异化贡献”的边界。更稳妥的表述是：围绕小病灶 CT-to-PET 合成，提出一个融合频率保护、医学先验和临床定量约束的 BBDM 框架。")
    doc.add_picture(str(figures["innovation"]), width=Inches(6.45))
    add_caption(doc, "图4  当前可向论文/答辩凝练的三个方法点：频率保护、先验轻量注入、代谢保真与不确定性。")
    innovation_rows = [
        ("尺度自适应布朗桥扩散", "Laplacian pyramid 分频加噪，高频噪声倍率低于低频，并用 Gabor energy 进一步保护纹理和边缘。", "需要通过 isotropic_noise / no_gabor 消融证明小病灶和边缘指标改善。"),
        ("多源先验的轻量条件注入", "CT encoder、Gabor、Hotspot、Organ/semantic 等条件经 Zero-conv adapter 注入 U-Net skip。", "需要 no_zero_adapter / no_hotspot / no_organ 消融，证明不是单纯参数增加。"),
        ("临床代谢保真与可置信生成", "Top-K lesion、ROI-SUV、false hotspot、heteroscedastic NLL、MC sampling 等已形成代码框架。", "PNG 当前无法完整支撑 SUV，下一步需 DICOM/NPZ cache 物理尺度闭环。"),
    ]
    t = doc.add_table(rows=1 + len(innovation_rows), cols=3)
    for j, h in enumerate(["方法点", "当前实现", "下一步证据"]):
        set_cell_text(t.cell(0, j), h, bold=True, color="FFFFFF", size=9.2)
        set_cell_shading(t.cell(0, j), NAVY)
        set_cell_border(t.cell(0, j), NAVY)
    for i, row in enumerate(innovation_rows, 1):
        for j, txt in enumerate(row):
            set_cell_text(t.cell(i, j), txt, bold=(j == 0), color=NAVY if j == 0 else "000000", size=8.8)
            if j == 0:
                set_cell_shading(t.cell(i, j), LIGHT_GRAY)
            set_cell_border(t.cell(i, j))
    doc.add_page_break()

    add_heading(doc, "6. 工程验证与当前证据边界", 1)
    add_body_para(doc, "在 ECPC 环境中补齐 `python-docx`、`pdf2image`、`reportlab`、`pytest` 等报告和测试依赖后，已运行当前项目 smoke tests：87 passed in 16.36s。测试覆盖了配置加载、消融开关、先验模块、条件注入、噪声日程、损失项、采样、MC 不确定性、数据集、split manifest、semantic builder 和训练器一步训练。")
    add_bullets(doc, [
        "已验证：模型能够前向计算 loss；DDIM/scale-adaptive bridge sampling 能返回 PET 预测；训练器可完成一步训练与梯度累积刷新。",
        "已验证：PNGSliceDataset 能按文件名配对 CT/PET/label，并保留原始 PET 尺寸供预测 PNG 回写。",
        "尚未验证：真实数据上的完整训练曲线、最佳 checkpoint、验证集 MAE/PSNR/SSIM/SUV 结果、患者级统计显著性。",
        "当前目录未发现 `checkpoints/`、`results/`、`runs/` 中的正式产物，因此答辩时建议把性能结果作为下一阶段重点，不提前放数值结论。",
    ])

    add_heading(doc, "7. 下一步工作计划", 1)
    doc.add_picture(str(figures["roadmap"]), width=Inches(6.45))
    add_caption(doc, "图5  中期后建议推进路线：先训练与验证，再做消融和临床定量扩展。")
    next_steps = [
        ("短期：跑通正式训练", "使用 `configs/experiments/slmf_png256.yaml` 跑 baseline 与 full，保存 checkpoint、resolved config、训练曲线和样例预测图。"),
        ("中期：完成量化评估", "在验证集输出 MAE/MSE/PSNR/SSIM、病灶 ROI 指标、false hotspot 粗略指标和患者级汇总表。"),
        ("中期：完成消融实验", "至少完成 baseline、no_gabor、no_hotspot、no_zero_adapter、isotropic_noise、no_heteroscedastic、no_self_conditioning、no_patch_nce。"),
        ("后期：回到 DICOM/NPZ 物理尺度", "构建 SUV 可用 cache，启用 ROI-SUV、organ consistency 和 false hotspot，使临床定量叙事更扎实。"),
        ("答辩材料输出", "形成训练曲线、预测可视化、失败案例分析和三项创新点对应的消融表格。"),
    ]
    t = doc.add_table(rows=1 + len(next_steps), cols=2)
    for j, h in enumerate(["任务", "具体动作"]):
        set_cell_text(t.cell(0, j), h, bold=True, color="FFFFFF", size=9.5)
        set_cell_shading(t.cell(0, j), NAVY)
        set_cell_border(t.cell(0, j), NAVY)
    for i, row in enumerate(next_steps, 1):
        for j, txt in enumerate(row):
            set_cell_text(t.cell(i, j), txt, bold=(j == 0), color=NAVY if j == 0 else "000000", size=9.0)
            if j == 0:
                set_cell_shading(t.cell(i, j), LIGHT_GRAY)
            set_cell_border(t.cell(i, j))

    add_heading(doc, "8. 可直接执行的关键命令", 1)
    commands = [
        "conda run -n ECPC python scripts/train_v2.py --config configs/experiments/slmf_png256.yaml",
        "conda run -n ECPC python scripts/train_v2.py --config configs/experiments/slmf_png256.yaml --ablation no_gabor",
        "conda run -n ECPC python scripts/evaluate.py --config configs/experiments/slmf_png256.yaml --checkpoint checkpoints/slmf_bbdm_png256/ckpt_epochXXXX.pt --split val --output results/slmf_png256_val.json --save-samples results/slmf_png256_val_samples",
        "conda run -n ECPC python -m pytest tests/test_smoke.py -q",
    ]
    for cmd in commands:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(cmd)
        set_run_font(r, name="Consolas", size=8.5, color=DARK_BLUE)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "附录：当前配置摘要", 1)
    add_body_para(doc, "PNG256 配置摘要：image_size=256，batch_size=4，gradient_accumulate_every=2，num_epochs=1000，learning_rate=1e-4，EMA=0.999，eval_sampling_steps=20。")
    add_body_para(doc, "Full 配置摘要：相同训练骨架下启用更多临床约束，包括 organ prior、ROI-SUV、false hotspot 和 organ consistency；segmenter consistency 仍建议在独立 segmenter 达到足够召回后再启用。")
    add_body_para(doc, f"消融预设数量：{len(ablations.get('ablations', {}))} 个。建议中期后优先跑与三项创新直接对应的核心消融。")

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)
