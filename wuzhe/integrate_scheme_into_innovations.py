from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SRC = Path("output/patent_disclosure/slmf_bbdm_patent_prior_art_optimized.docx")
OUT = Path("output/patent_disclosure/slmf_bbdm_patent_innovation_integrated.docx")


def set_east_asian_font(run, font="SimSun") -> None:
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font)


def insert_paragraph_after(paragraph, text: str, style_name: str = "Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    # Move the newly-created paragraph XML to the intended position.
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    paragraph._p.getparent().remove(new_p)
    new_para.style = style_name
    run = new_para.add_run(text)
    set_east_asian_font(run)
    # Match the prevailing body rhythm where possible.
    new_para.paragraph_format.space_after = paragraph.paragraph_format.space_after
    new_para.paragraph_format.line_spacing = paragraph.paragraph_format.line_spacing
    return new_para


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        for run in paragraph.runs:
            run.text = ""
        first.text = text
        set_east_asian_font(first)
    else:
        run = paragraph.add_run(text)
        set_east_asian_font(run)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def main() -> None:
    doc = Document(SRC)

    # Integrate the former "technical scheme" into the two innovation points.
    innovation_one_additions = [
        "在该创新点的具体实施中，模型首先对同一患者、同一或近似层面的 CT 与 PET 图像进行配准质控和切片匹配；对于 DICOM 数据，读取 CT 像素并转换为 HU 值，将 HU 裁剪到预设窗口后归一化至 [-1,1]，同时在具备体重、注射剂量、采集时间和放射性核素半衰期等字段时计算 PET SUV 并归一化。上述输入构建不是单独的数据预处理步骤，而是频带桥式状态构造的前置限定，使 CT 端点、PET 端点、SUV 标定和器官分区在同一坐标与强度尺度下参与后续噪声门控。",
        "具体地，可采用 CT_norm = 2·clip((HU-HU_min)/(HU_max-HU_min),0,1)-1 以及 PET_norm = 2·clip(SUV/SUV_max,0,1)-1 进行归一化；同时依据 CT HU 构造 511keV 衰减映射 μ_map，并将子宫或盆腔区域、膀胱、直肠或肠道、骨、脂肪、肌肉或其他组织映射为器官分区通道。该器官分区不作为普通分割结果单独使用，而是与方向频率响应、热点候选概率共同决定 ρ_b(p)、R_cold(p) 及后续临床损失是否启用。",
        "对于 PNG 或其他二维图像格式，也可采用文件名匹配方式配对 CT、PET 和标签图，并统一缩放到模型输入大小；若需要严格 SUV 损失和物理尺度分析，则优选 DICOM 到 NPZ 缓存流程。由此，图像输入、物理标定和配准质控被并入创新点一的桥式状态构造链条，而不是作为独立技术方案列出。",
    ]

    # Insert after the last paragraph of innovation point one.
    anchor_one = doc.paragraphs[29]
    for text in reversed(innovation_one_additions):
        insert_paragraph_after(anchor_one, text)

    innovation_two_additions = [
        "在该创新点的具体实施中，CT 多尺度编码器用于提取不同感受野的解剖结构特征，方向频率响应模块用于生成高频能量图，热点候选模块用于输出病灶邻域概率图，器官先验模块用于把器官类别转化为代谢规则。上述模块不再作为单独技术方案分项列出，而是作为时间门控轻量适配机制的条件来源：浅层尺度优先接收 CT 细节、方向频率响应和热点候选图，深层尺度优先接收器官分区、器官距离、衰减映射和病例元数据。",
        "优选地，方向频率响应能量可由 Gabor、小波、LoG、Sobel、多方向可学习卷积或频域滤波器产生；热点候选图由病灶掩膜、PET 高摄取分位区域和可选距离图共同监督；器官分区中膀胱和直肠或肠道被视为允许生理摄取区域，骨、脂肪和肌肉被视为代谢冷区，子宫或盆腔区域采用弱约束或候选增强约束。这样，编码、先验和去噪网络都服从创新点二的“器官代谢规则驱动的时间门控适配”主线。",
        "去噪主干可采用残差 U-Net 或其他编码器—解码器去噪网络，其输入为当前桥式扩散状态 x_t 与源 CT x_s 的拼接，输出为预测 PET 图像和可选 log-variance 通道。其保护重点不在去噪主干本身，而在于其接收的桥式状态、浅深层条件适配、器官代谢规则和阶段化 SUV/假热点损失之间的配合关系。",
        "在公式表达上，浅层条件可写为 C_0=concat(f_0^CT,f^dir,A^hotspot)，深层条件可写为 C_l=concat(f_l^CT,f_l^organ)，l∈{1,2,3}，尺度适配量为 Δh_l=β_l(τ)·Conv1×1_zero(C_l)。由于适配单元在初始化时为零或近零，模型初始状态等价于未注入先验的基础去噪主干，训练后再由时间门控逐步学习各类先验贡献。",
    ]

    anchor_two = doc.paragraphs[37]
    for text in reversed(innovation_two_additions):
        insert_paragraph_after(anchor_two, text)

    # Move figure 4 (former PET/CT QC figure) under innovation point one.
    # In the source it appears as paragraphs 51-52. After insertions, locate by caption.
    fig4_caption_idx = None
    for i, p in enumerate(doc.paragraphs):
        if "图4  PET/CT 数据配准与质控流程示意图" in p.text:
            fig4_caption_idx = i
            break
    if fig4_caption_idx is not None and fig4_caption_idx > 0:
        fig_p = doc.paragraphs[fig4_caption_idx - 1]
        cap_p = doc.paragraphs[fig4_caption_idx]
        fig_xml = fig_p._element
        cap_xml = cap_p._element
        fig_xml.getparent().remove(fig_xml)
        cap_xml.getparent().remove(cap_xml)
        # Place it after the third inserted innovation-one paragraph.
        target = doc.paragraphs[32]._element
        target.addnext(fig_xml)
        fig_xml.addnext(cap_xml)

    # Delete the standalone "技术方案" section and its six sub-sections.
    # Remove from heading "技术方案" up to before "7.损失函数与公式化描述".
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "技术方案" and p.style.name.startswith("Heading"):
            start = i
        if start is not None and txt.startswith("7.损失函数"):
            end = i
            break
    if start is not None and end is not None:
        for p in list(doc.paragraphs[start:end]):
            remove_paragraph(p)

    # Since the standalone technical scheme is gone, remove the orphaned number.
    for p in doc.paragraphs:
        if p.text.strip().startswith("7.损失函数"):
            replace_paragraph_text(p, "损失函数与公式化描述")
            break

    # Renumber figure captions after moving the former figure 4 into innovation point one.
    caption_replacements = {
        "图4  PET/CT 数据配准与质控流程示意图": "图1  PET/CT 数据配准与质控流程示意图",
        "图1  SLMF-BBDM 总体网络结构示意图": "图2  SLMF-BBDM 总体网络结构示意图",
        "图2  CT-to-PET 布朗桥扩散训练与推理流程图": "图3  CT-to-PET 布朗桥扩散训练与推理流程图",
        "图3  两项核心创新点与技术闭环示意图": "图4  两项核心创新点与技术闭环示意图",
    }
    figure_notes = {
        "图1 为 SLMF-BBDM 总体网络结构示意图，展示 CT 编码、方向频率响应、热点候选、盆腔器官分区、时间门控尺度适配、桥式扩散和去噪网络之间的限定关系。":
            "图1 为 PET/CT 数据配准与质控流程示意图，展示模型训练前的数据一致性检查，并说明该质控流程已融入创新点一的桥式状态构造前置条件。",
        "图2 为训练与推理流程图，展示真实 PET、源 CT、前向桥式加噪、条件构建、反向采样和合成 PET 输出。":
            "图2 为 SLMF-BBDM 总体网络结构示意图，展示 CT 编码、方向频率响应、热点候选、盆腔器官分区、时间门控尺度适配、桥式扩散和去噪网络之间的限定关系。",
        "图3 为核心创新点示意图，展示频带桥式状态构造、盆腔器官代谢规则注入和 SUV 阶段化误差重分配之间的关系。":
            "图3 为 CT-to-PET 布朗桥扩散训练与推理流程图，展示真实 PET、源 CT、前向桥式加噪、条件构建、反向采样和合成 PET 输出。",
        "图4 为 PET/CT 数据配准与质控流程示意图，展示模型训练前的数据一致性检查。":
            "图4 为核心创新点示意图，展示频带桥式状态构造、盆腔器官代谢规则注入和 SUV 阶段化误差重分配之间的关系。",
    }
    all_replacements = {**caption_replacements, **figure_notes}
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in all_replacements:
            replace_paragraph_text(p, all_replacements[txt])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
